"""Every instruction the models receive, as one Word document.

Shared by the download button and `scripts/export_prompts.py`, because a
document generated two ways drifts into two documents.

What it includes, and why all of it:

* **The preamble**, which is not editable anywhere. Leaving it out would make
  the document describe something the models are never actually sent — every
  stage runs as preamble + body, and a reviewer reading only the bodies would
  be reviewing two thirds of the instruction.
* **The live body** of all sixteen stages, read from the database, so an edit
  made in the UI is in the next export.
* **The shipped original**, but only where someone has changed it. Printing it
  for every stage would double the length to say "unchanged" fifteen times.

Stages come out in the order `instructions.STAGE` declares them, which is
roughly the order the pipeline runs — research, then extraction, then
qualification, strategy and content. Alphabetical would scatter related stages,
and is what the screen already does.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from sqlalchemy.ext.asyncio import AsyncSession

from app.services import instructions, prompts

MONO = "Consolas"
MUTED = RGBColor(0x5A, 0x62, 0x6A)
WARN = RGBColor(0x9A, 0x5B, 0x00)

FILENAME = "prompt-reference.docx"
MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@dataclass
class ExportedStage:
    scenario: str
    body: str
    version: int | None
    note: str | None
    edited: bool
    shipped: str
    stored: bool


async def collect(session: AsyncSession) -> list[ExportedStage]:
    live = await prompts.active_map(session)
    out: list[ExportedStage] = []
    for scenario in instructions.STAGE:
        row = live.get(scenario)
        default = prompts.shipped_default(scenario)
        out.append(
            ExportedStage(
                scenario=scenario,
                body=row.body if row else default,
                version=row.version if row else None,
                note=row.note if row else None,
                edited=bool(row and row.body.strip() != default),
                shipped=default,
                stored=row is not None,
            )
        )
    return out


def _mono(doc: Document, text: str, *, shaded: bool = False) -> None:
    """A prompt, rendered as what it is: text sent verbatim to a model.

    Monospaced and kept line for line rather than reflowed into prose. The line
    breaks are load-bearing — several blocks are lists the model reads item by
    item — and Word's default paragraph spacing turns them into something that
    reads differently from what is actually sent.
    """
    for line in text.split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Pt(14)
        run = p.add_run(line or " ")
        run.font.name = MONO
        run.font.size = Pt(9)
        if shaded:
            run.font.color.rgb = MUTED


def _note(doc: Document, text: str, *, colour: RGBColor | None = None) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(9)
    run.italic = True
    run.font.color.rgb = colour or MUTED


def build(stages: list[ExportedStage], *, today: date | None = None) -> bytes:
    """The document, as bytes — written to a file or streamed, not both."""
    doc = Document()

    title = doc.add_heading("Prospect Engine — model instructions", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _note(
        doc,
        "Every instruction the language models receive. Exported "
        f"{(today or date.today()):%d %B %Y} from the live configuration.",
    )

    doc.add_heading("How to read this", level=1)
    doc.add_paragraph(
        "Each stage is sent two things, in this order: the common preamble, "
        "then that stage's own instructions. The preamble is identical for "
        "every stage and is reproduced once, below. Nothing else is added — "
        "what follows is the complete text, verbatim."
    )
    doc.add_paragraph(
        "The stage instructions are editable in the application, under "
        "Settings → Prompts, by anyone holding the admin password. Each save "
        "creates a new version and the previous one is kept, so the output of "
        "any past run can still be traced to the instructions that produced it."
    )
    doc.add_paragraph(
        "The preamble is not editable, there or anywhere. It carries the rule "
        "that retrieved text is data rather than instruction — the guard "
        "against a researched web page telling the model what to do — and the "
        "rule that nothing is asserted without a citation. Those are "
        "properties of the system rather than settings within it."
    )

    edited = [s for s in stages if s.edited]
    doc.add_paragraph(
        f"{len(stages)} stages. "
        + (
            f"{len(edited)} edited from the shipped text: "
            + ", ".join(s.scenario for s in edited)
            + "."
            if edited
            else "None has been edited from the shipped text."
        )
    )

    doc.add_page_break()
    doc.add_heading("The common preamble", level=1)
    _note(
        doc,
        "Prepended to every stage below. Not editable in the application.",
        colour=WARN,
    )
    doc.add_paragraph()
    _mono(doc, instructions.COMMON.strip())

    doc.add_page_break()
    doc.add_heading("Stage instructions", level=1)

    for i, s in enumerate(stages):
        if i:
            doc.add_page_break()
        doc.add_heading(s.scenario, level=2)

        bits = [f"version {s.version}" if s.version else "no stored version"]
        bits.append("edited" if s.edited else "unchanged from shipped")
        bits.append(f"{len(s.body)} characters")
        _note(doc, " · ".join(bits))
        if s.note:
            _note(doc, f"Reason given for the current version: {s.note}")
        if not s.stored:
            _note(
                doc,
                "No stored row — this stage is running on its shipped text.",
                colour=WARN,
            )
        doc.add_paragraph()
        _mono(doc, s.body.strip())

        if s.edited:
            doc.add_paragraph()
            doc.add_heading("Shipped original, for comparison", level=3)
            _note(doc, "What this stage said before it was edited here.")
            doc.add_paragraph()
            _mono(doc, s.shipped.strip(), shaded=True)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


async def export(session: AsyncSession) -> bytes:
    return build(await collect(session))
